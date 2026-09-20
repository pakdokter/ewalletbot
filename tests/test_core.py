import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.excel_export import build_workbook
from app.models import Txn
from app.ocr.base import OcrPage, Word
from app.parser import is_main_balance, parse_date_text, parse_page
from app.session import Session, merge_overlap, parse_date_input, parse_rupiah


# ---------- angka & tanggal
@pytest.mark.parametrize("text,expected", [
    ("62620.36", 62620.36), ("1.250.000", 1250000), ("Rp 500.000", 500000),
    ("1.250,50", 1250.5), ("-114300", -114300), ("2500", 2500),
])
def test_parse_rupiah(text, expected):
    assert parse_rupiah(text) == pytest.approx(expected)


def test_parse_date_input():
    assert parse_date_input("31/01/2024") == "2024-01-31"
    assert parse_date_input("2024-01-31") == "2024-01-31"
    with pytest.raises(ValueError):
        parse_date_input("besok")


def test_date_header_variants_and_weekday_check():
    d, ok = parse_date_text("Senin 1Jan 2024")           # spasi hilang seperti pada OCR nyata
    assert d == date(2024, 1, 1) and ok is True
    d, ok = parse_date_text("Selasa 2 Jan 2024")
    assert d == date(2024, 1, 2) and ok is True
    d, ok = parse_date_text("Rabu 2 Jan 2024")           # hari salah -> terdeteksi
    assert ok is False
    assert parse_date_text("12 Agu 2024")[0] == date(2024, 8, 12)
    assert parse_date_text("tidak ada tanggal") is None


def test_main_balance_split():
    assert is_main_balance("GoPay Saldo", "GoPay")
    assert not is_main_balance("GoPay Coins", "GoPay")
    assert not is_main_balance("GoPay Later", "GoPay")
    assert is_main_balance("OVO Cash", "OVO")
    assert not is_main_balance("OVO Points", "OVO")


# ---------- tumpang tindih
def T(desc, amt, d="2024-01-01"):
    return Txn(d, desc, amt, "GoPay Saldo", "GoPay")


def test_overlap_only_at_boundary():
    prev = [T("a", -1), T("b", -2), T("c", -3)]
    new = [T("b", -2), T("c", -3), T("d", -4)]
    fresh, dropped = merge_overlap(prev, new)
    assert dropped == 2 and [t.desc for t in fresh] == ["d"]


def test_identical_rows_inside_one_screenshot_are_kept():
    prev = [T("a", -1)]
    new = [T("kopi", -24000), T("kopi", -24000)]          # dua pembelian identik, bukan duplikat
    fresh, dropped = merge_overlap(prev, new)
    assert dropped == 0 and len(fresh) == 2


def test_chronological_order_reverses_within_day():
    s = Session(1)
    s.add_rows([T("siang", -1, "2024-01-02"), T("sore", -2, "2024-01-01"), T("pagi", -3, "2024-01-01")])
    # layar: baris atas = terbaru. Dalam 1 Jan: "sore" (atas) lebih baru dari "pagi"
    assert [t.desc for t in s.chronological()] == ["pagi", "sore", "siang"]


# ---------- parser pada tata letak sintetis (tanpa OCR)
def W(text, x, y, h=40, conf=95, w=None):
    return Word(text, x, y, w or 14 * len(text), h, conf)


def synthetic_page():
    words = [
        W("Riwayat", 100, 60), W("transaksi", 300, 60), W("Layanan", 800, 200), W("Metode", 1100, 200),
        W("Januari", 200, 340), W("2024", 420, 340),
        W("Selasa", 200, 460), W("2", 380, 460), W("Jan", 420, 460), W("2024", 520, 460),
        W("Senin", 200, 620), W("1Jan", 380, 620), W("2024", 520, 620),
        # baris 1: debit
        W("lapakgaming", 400, 780), W("-Rp114.300", 1150, 760, h=44),
        W("GoPay", 1180, 830), W("Saldo", 1320, 830),
        # baris 2: kredit koin (tanpa Rp)
        W("Cashback", 400, 1080), W("MyTelkomsel", 620, 1080), W("2.000", 1380, 1060, h=44),
        W("GoPay", 1180, 1130), W("Coins", 1320, 1130),
        # baris 3: kredit saldo
        W("GoPay", 400, 1380), W("Top", 540, 1380), W("Up", 620, 1380), W("Rp500.000", 1250, 1360, h=44),
        W("GoPay", 1180, 1430), W("Saldo", 1320, 1430),
        W("Beranda", 100, 1900), W("Keuangan", 300, 1900), W("Riwayat", 900, 1900), W("Profil", 1300, 1900),
    ]
    return OcrPage(words, 1600, 2000, None, False, "test", None)


def test_parser_reads_layout():
    res = parse_page(synthetic_page())
    assert res.wallet == "GoPay"
    got = [(t.date, t.desc, t.amount, t.method) for t in res.rows]
    assert got == [
        ("2024-01-01", "lapakgaming", -114300, "GoPay Saldo"),
        ("2024-01-01", "Cashback MyTelkomsel", 2000, "GoPay Coins"),
        ("2024-01-01", "GoPay Top Up", 500000, "GoPay Saldo"),
    ]
    assert all(not t.flags for t in res.rows)


def test_parser_flags_low_confidence_amount_and_missing_date():
    page = synthetic_page()
    for w in page.words:
        if w.text == "-Rp114.300":
            w.conf = 40
        if w.text in {"Senin", "1Jan"}:
            w.text = "xx"                                   # tanggal rusak
    res = parse_page(page, carry_date=None)
    assert any("nominal kurang yakin" in f for f in res.rows[0].flags)
    assert any("tanggal meragukan" in f for f in res.rows[0].flags)   # tidak boleh diam-diam memakai tanggal 2 Jan


# ---------- Excel
def test_excel_export_structure_and_formulas(tmp_path):
    s = Session(1, wallet="GoPay", saldo_awal=62620.36)
    s.add_rows([
        Txn("2024-01-01", "GoPay Top Up", 500000, "GoPay Saldo", "GoPay"),
        Txn("2024-01-01", "MyTelkomsel", -146000, "GoPay Saldo", "GoPay"),
        Txn("2024-01-01", "Cashback MyTelkomsel", 2000, "GoPay Coins", "GoPay"),
        Txn("2024-01-01", "lapakgaming", -114300, "GoPay Saldo", "GoPay", ["nama kurang jelas"]),
    ])
    out = tmp_path / "x.xlsx"
    info = build_workbook(s, out)
    assert info["main"] == 3 and info["other"] == 1 and info["flagged"] == 1

    wb = load_workbook(out)
    ws = wb["Mutasi"]
    assert [c.value for c in ws[1]] == ["Tanggal", "Keterangan Transaksi", "Kategori Transaksi", "Debit", "Kredit",
                                        "Saldo Kumulatif", "Subjek Transaksi", "Objek Transaksi", "Keterangan Tambahan"]
    assert ws["B2"].value == "Saldo Awal" and ws["F2"].value == pytest.approx(62620.36)
    # kronologis dalam satu hari: urutan layar dibalik (layar: TopUp, MyTelkomsel, lapakgaming = terbaru -> lama? seq kecil = atas)
    assert [ws.cell(r, 2).value for r in range(3, 6)] == ["lapakgaming", "MyTelkomsel", "GoPay Top Up"]
    assert ws["F3"].value == "=D3+E3+F2"
    assert ws["C3"].value is None and ws["I3"].value is None            # kolom manual dibiarkan kosong
    # konvensi Subjek/Objek: debit -> (dompet, merchant); kredit -> (sumber, DOMPET)
    assert (ws["G3"].value, ws["H3"].value) == ("GoPay", "lapakgaming")
    assert (ws["G5"].value, ws["H5"].value) == ("GoPay Top Up", "GOPAY")
    assert ws["B3"].comment is not None                                   # baris bertanda diberi komentar
    assert "Non-Saldo" in wb.sheetnames and "Catatan" in wb.sheetnames
    assert wb["Non-Saldo"]["B2"].value == "Cashback MyTelkomsel [GoPay Coins]"
    assert ws["D8"].value == "=SUM(D3:D5)" and ws["E9"].value == "=SUM(E3:E5)"
