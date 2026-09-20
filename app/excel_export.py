"""Buat file Excel dengan format acuan Jago_September_2024.xlsx.

Kolom:  A Tanggal | B Keterangan Transaksi | C Kategori Transaksi (MANUAL) | D Debit | E Kredit |
        F Saldo Kumulatif (rumus) | G Subjek Transaksi | H Objek Transaksi | I Keterangan Tambahan (MANUAL)
Konvensi G/H mengikuti acuan:  debit  -> G = dompet, H = merchant
                                kredit -> G = sumber,  H = DOMPET (huruf besar)
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from .models import Txn
from .parser import is_main_balance
from .session import Session

FONT = "Arial"
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADERS = ["Tanggal", "Keterangan Transaksi", "Kategori Transaksi", "Debit", "Kredit", "Saldo Kumulatif",
           "Subjek Transaksi", "Objek Transaksi", "Keterangan Tambahan"]
WIDTHS = {"A": 12, "B": 24, "C": 24, "D": 16, "E": 16, "F": 16, "G": 20, "H": 24, "I": 40}   # dari acuan

FMT_DEBIT = r'"Rp"\ #,##0.00;\("Rp"\ #,##0.00\)'
FMT_MONEY = r'"Rp"\ #,##0.00'
FMT_DATE = "dd/mm/yyyy"

FILL_HEADER = "1F5C22"
FILL_OPENING = "FDF0C7"
FILL_BALANCE_COL = "FCD9B6"
FILL_SUMMARY = "D9EAD9"

# Warna baris per kategori, diambil dari file acuan. Kolom C diisi manual -> warna muncul otomatis
# lewat conditional formatting begitu kategori diketik. Tambahkan kategori barumu di sini.
CATEGORY_FILLS = {
    "Transaksi Internal": "FFF2A8",
    "Konsumsi dan Liburan": "F7C9DE",
    "Marketing": "F7C9DE",
    "Gaji Accrual": "FFCC99",
    "Gaji Bulan Ini": "FFCC99",
    "Belanja Bahan": "E4C9A0",
    "Subscription": "D3D3D3",
    "Pengeluaran Pribadi": "E8AB8C",
    "Biaya Admin Bank": "D9D9A3",
}


def _solid(hex_: str) -> PatternFill:
    return PatternFill(start_color=hex_, end_color=hex_, fill_type="solid")


def _font(bold: bool = False, color: str | None = None) -> Font:
    return Font(name=FONT, size=9, bold=bold, color=color)


def _header(ws: Worksheet) -> None:
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(1, i, h)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = _solid(FILL_HEADER)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    ws.row_dimensions[1].height = 18
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def _as_date(iso: str | None):
    return date.fromisoformat(iso) if iso else None


def _parties(t: Txn) -> tuple[str, str]:
    """(Subjek, Objek) sesuai konvensi acuan."""
    if t.amount < 0:
        return t.wallet, t.desc
    return t.desc, t.wallet.upper()


def _write_row(ws: Worksheet, r: int, t: Txn, *, chain_balance: bool, label: str | None = None) -> None:
    g, h = _parties(t)
    vals = {1: _as_date(t.date), 2: label or t.desc, 3: None,
            4: t.amount if t.amount < 0 else None, 5: t.amount if t.amount > 0 else None,
            6: f"=D{r}+E{r}+F{r-1}" if chain_balance else "-", 7: g, 8: h, 9: None}
    for col, v in vals.items():
        c = ws.cell(r, col, v)
        c.font = _font()
        c.border = BORDER
    ws.cell(r, 1).number_format = FMT_DATE
    ws.cell(r, 4).number_format = FMT_DEBIT
    ws.cell(r, 5).number_format = FMT_MONEY
    ws.cell(r, 6).number_format = FMT_MONEY
    ws.cell(r, 6).fill = _solid(FILL_BALANCE_COL)
    if t.flags:
        b = ws.cell(r, 2)
        b.font = _font(color="C00000")
        b.comment = Comment("Perlu dicek (OCR kurang yakin): " + "; ".join(t.flags), "Bot")


def build_workbook(sess: Session, path: Path) -> dict:
    """Tulis xlsx. Mengembalikan ringkasan (jumlah baris dst.) untuk pesan balasan."""
    rows = sess.chronological()
    wallet = sess.wallet or (rows[0].wallet if rows else "Dompet")
    main = [t for t in rows if is_main_balance(t.method, wallet)]
    other = [t for t in rows if not is_main_balance(t.method, wallet)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Mutasi"
    _header(ws)

    # baris saldo awal (mengikuti acuan: nilai di kolom E dan F)
    opening = float(sess.saldo_awal or 0)
    for col, v in {2: "Saldo Awal", 3: "Saldo Awal Bulan", 5: opening, 6: opening, 7: "-", 8: "-"}.items():
        ws.cell(2, col, v)
    for col in range(1, 10):
        c = ws.cell(2, col)
        c.font, c.fill, c.border = _font(bold=True), _solid(FILL_OPENING), BORDER
    ws.cell(2, 5).number_format = FMT_MONEY
    ws.cell(2, 6).number_format = FMT_MONEY
    if sess.saldo_awal is None:
        ws.cell(2, 6).comment = Comment("Saldo awal belum diisi -> dianggap 0. Ubah sel E2 dan F2, "
                                        "atau kirim /saldoawal lalu /rekap lagi.", "Bot")

    r = 3
    for t in main:
        _write_row(ws, r, t, chain_balance=True)
        r += 1
    last = r - 1

    # warna baris berdasarkan kategori (kolom C manual)
    if main:
        for cat, hex_ in CATEGORY_FILLS.items():
            ws.conditional_formatting.add(
                f"A3:E{last} G3:I{last}",
                FormulaRule(formula=[f'$C3="{cat}"'], fill=_solid(hex_)))
        ws.auto_filter.ref = f"A1:I{last}"

    # ringkasan bawah - rumus, bukan angka mati
    s = max(last, 2) + 2
    end = max(last, 3)
    summary = [
        ("Saldo Awal", 6, "=F2"),
        ("Total Debit (Uang Keluar)", 4, f"=SUM(D3:D{end})"),
        ("Total Kredit (Uang Masuk)", 5, f"=SUM(E3:E{end})"),
        ("Saldo Akhir", 6, f"=F{s}+D{s+1}+E{s+2}"),
    ]
    for i, (label, col, formula) in enumerate(summary):
        rr = s + i
        for cc in range(1, 10):
            c = ws.cell(rr, cc)
            c.font, c.fill, c.border = _font(bold=True), _solid(FILL_SUMMARY), BORDER
        ws.cell(rr, 2, label)
        v = ws.cell(rr, col, formula)
        v.number_format = FMT_DEBIT if col == 4 else FMT_MONEY

    # baris non-saldo (Coins / Points / Later ...) - tidak memengaruhi saldo dompet
    if other:
        wo = wb.create_sheet("Non-Saldo")
        _header(wo)
        for i, t in enumerate(other, start=2):
            _write_row(wo, i, t, chain_balance=False, label=f"{t.desc} [{t.method}]")

    # catatan / legenda
    wn = wb.create_sheet("Catatan")
    dates = [t.date for t in rows if t.date]
    flagged = sum(1 for t in rows if t.flags)
    lines = [
        (f"Rekap transaksi {wallet} - disusun otomatis dari screenshot oleh bot.", True),
        (f"BUKAN e-statement resmi {wallet}. Untuk rekonsiliasi dan analisis internal.", True),
        ("", False),
        (f"Dibuat: {datetime.now():%d/%m/%Y %H:%M}", False),
        (f"Periode data: {min(dates) if dates else '-'} s.d. {max(dates) if dates else '-'}", False),
        (f"Baris di sheet Mutasi: {len(main)} | Non-Saldo: {len(other)} | Ditandai perlu dicek: {flagged}", False),
        ("", False),
        ("Cara pakai", True),
        ("- Kolom C (Kategori Transaksi) dan I (Keterangan Tambahan) diisi manual. Contoh C: Belanja Bahan; contoh I: Paid to Gopay Owner.", False),
        ("- Warna baris muncul otomatis begitu kategori dari file acuan diketik di kolom C (daftar ada di CATEGORY_FILLS pada excel_export.py).", False),
        ("- Debit bertanda minus, Kredit positif. Saldo Kumulatif dan ringkasan bawah adalah rumus.", False),
        ("- Nama merchant berhuruf merah + komentar = OCR kurang yakin. Cocokkan dengan layar aplikasi, lalu perbaiki.", False),
        ("", False),
        ("Asumsi yang perlu kamu ketahui", True),
        ("- Layar aplikasi menampilkan transaksi terbaru di atas; urutan dalam satu hari dibalik supaya kronologis. Saldo antar-hari pasti benar, saldo di tengah hari bergantung asumsi ini.", False),
        ("- Saldo awal berasal dari perintah /saldoawal (bila kosong dianggap 0). Screenshot riwayat tidak memuat saldo.", False),
        ("- Coins / Points / Later bukan saldo dompet, jadi dipisah ke sheet Non-Saldo.", False),
        ("- Kontrol utama: samakan 'Saldo Akhir' dengan saldo dompet di aplikasi pada tanggal terakhir. Selisih = ada baris yang salah baca atau terlewat.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = wn.cell(i, 1, text)
        c.font = _font(bold=bold)
    wn.column_dimensions["A"].width = 140

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"main": len(main), "other": len(other), "flagged": flagged,
            "first": min(dates) if dates else None, "last": max(dates) if dates else None}
